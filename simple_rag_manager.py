# simple_rag_manager.py
import os
import logging
import uuid
import hashlib
import json
import sqlite3
from typing import List, Dict, Any, Optional, Tuple, Union
from pathlib import Path
import re
from sentence_transformers import SentenceTransformer
import numpy as np
import torch
import pickle
import scipy.sparse
# Add PyMuPDF for PDF processing
import pymupdf  # PyMuPDF
import docx  # For DOCX files

from PyQt5.QtCore import QThread, QObject, pyqtSignal

# Simple vectorization using scikit-learn (much lighter than transformer models)
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# Set up logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class Document:
    """Represents a document with metadata and content"""
    
    def __init__(
        self, 
        doc_id: str,
        filename: str,
        content: str,
        metadata: Optional[Dict[str, Any]] = None,
        chunks: Optional[List[str]] = None
    ):
        self.doc_id = doc_id
        self.filename = filename
        self.content = content
        self.metadata = metadata or {}
        self.chunks = chunks or []
        
    def to_dict(self) -> Dict[str, Any]:
        """Convert document to dictionary for storage"""
        return {
            "id": self.doc_id,
            "filename": self.filename,
            "metadata": self.metadata,
            "content": self.content
        }

class DocumentChunker:
    """Splits documents into manageable chunks"""
    
    def __init__(self, chunk_size: int = 1500, chunk_overlap: int = 500):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
    
    def chunk_document(self, document: Document) -> Document:
        """Split document content into overlapping chunks"""
        content = document.content
        chunks = []
        
        # Simple chunking by characters with overlap
        for i in range(0, len(content), self.chunk_size - self.chunk_overlap):
            chunk = content[i:i + self.chunk_size]
            if len(chunk) >= 100:  # Only keep chunks with sufficient content
                chunks.append(chunk)
        
        document.chunks = chunks
        return document

class DocumentProcessor:
    """Processes different document types into text"""
    
    @staticmethod
    def process_file(file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Process file into text content and metadata"""
        file_ext = os.path.splitext(file_path)[1].lower()
        
        # Handle DOCX files
        if file_ext == '.docx':
            try:
                # Process DOCX with python-docx
                doc = docx.Document(file_path)
                
                # Extract text from paragraphs
                text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                
                # DOCX-specific metadata
                metadata = {
                    "format": "docx",
                    "size_bytes": os.path.getsize(file_path),
                    "paragraph_count": len(doc.paragraphs),
                    "line_count": text.count('\n') + 1
                }
                
                return text, metadata
                
            except Exception as e:
                logger.error(f"Error processing DOCX file: {str(e)}")
                return "", {"error": f"DOCX processing error: {str(e)}"}
        
        # Handle PDF files
        if file_ext == '.pdf':
            try:
                # Process PDF with PyMuPDF
                doc = pymupdf.open(file_path)
                text = ""
                page_count = 0
                
                # Extract text from each page
                for page_num in range(len(doc)):
                    page = doc[page_num]
                    text += page.get_text()
                    page_count += 1
                
                doc.close()
                
                # PDF-specific metadata
                metadata = {
                    "format": "pdf",
                    "size_bytes": os.path.getsize(file_path),
                    "page_count": page_count,
                    "line_count": text.count('\n') + 1
                }
                
                return text, metadata
                
            except Exception as e:
                logger.error(f"Error processing PDF file: {str(e)}")
                return "", {"error": f"PDF processing error: {str(e)}"}
        
        # Handle PowerPoint files
        if file_ext == '.pptx':
            try:
                from pptx import Presentation
                prs = Presentation(file_path)
                text_parts = []
                slide_count = 0
                for slide in prs.slides:
                    slide_count += 1
                    slide_text = f"--- Slide {slide_count} ---\n"
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                slide_text += paragraph.text + "\n"
                        if shape.has_table:
                            for row in shape.table.rows:
                                row_text = " | ".join(cell.text for cell in row.cells)
                                slide_text += row_text + "\n"
                    text_parts.append(slide_text)
                text = "\n".join(text_parts)
                metadata = {
                    "format": "pptx",
                    "size_bytes": os.path.getsize(file_path),
                    "slide_count": slide_count,
                    "line_count": text.count('\n') + 1
                }
                return text, metadata
            except ImportError:
                logger.error("python-pptx not installed. Install with: pip install python-pptx")
                return "", {"error": "python-pptx not installed. Run: pip install python-pptx"}
            except Exception as e:
                logger.error(f"Error processing PPTX file: {str(e)}")
                return "", {"error": f"PPTX processing error: {str(e)}"}

        # Handle Excel files
        if file_ext in ('.xlsx', '.xls'):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                text_parts = []
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    text_parts.append(f"--- Sheet: {sheet_name} ---")
                    for row in ws.iter_rows(values_only=True):
                        row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                        if row_text.strip():
                            text_parts.append(row_text)
                wb.close()
                text = "\n".join(text_parts)
                metadata = {
                    "format": file_ext.lstrip('.'),
                    "size_bytes": os.path.getsize(file_path),
                    "sheet_count": len(wb.sheetnames),
                    "line_count": text.count('\n') + 1
                }
                return text, metadata
            except ImportError:
                logger.error("openpyxl not installed. Install with: pip install openpyxl")
                return "", {"error": "openpyxl not installed. Run: pip install openpyxl"}
            except Exception as e:
                logger.error(f"Error processing Excel file: {str(e)}")
                return "", {"error": f"Excel processing error: {str(e)}"}

        # Handle image files with OCR
        if file_ext in ('.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif'):
            try:
                # Try PyMuPDF first for image text extraction
                doc = pymupdf.open(file_path)
                page = doc[0]
                text = page.get_text()
                doc.close()

                if not text.strip():
                    # Fallback: try pytesseract OCR
                    try:
                        from PIL import Image
                        import pytesseract
                        img = Image.open(file_path)
                        text = pytesseract.image_to_string(img)
                    except ImportError:
                        text = f"[Image file: {os.path.basename(file_path)} - OCR not available. Install Pillow and pytesseract for image text extraction.]"
                    except Exception as ocr_err:
                        text = f"[Image file: {os.path.basename(file_path)} - OCR failed: {str(ocr_err)}]"

                metadata = {
                    "format": file_ext.lstrip('.'),
                    "size_bytes": os.path.getsize(file_path),
                    "type": "image",
                    "line_count": text.count('\n') + 1
                }
                return text, metadata
            except Exception as e:
                logger.error(f"Error processing image file: {str(e)}")
                return "", {"error": f"Image processing error: {str(e)}"}

        # Handle CSV files
        if file_ext == '.csv':
            try:
                import csv
                text_parts = []
                with open(file_path, 'r', encoding='utf-8', newline='') as f:
                    reader = csv.reader(f)
                    for row in reader:
                        text_parts.append(" | ".join(row))
                text = "\n".join(text_parts)
                metadata = {
                    "format": "csv",
                    "size_bytes": os.path.getsize(file_path),
                    "line_count": len(text_parts)
                }
                return text, metadata
            except Exception as e:
                logger.error(f"Error processing CSV file: {str(e)}")
                return "", {"error": f"CSV processing error: {str(e)}"}

        # For text files and other formats
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()

            metadata = {
                "format": file_ext.lstrip('.') if file_ext else "txt",
                "size_bytes": os.path.getsize(file_path),
                "line_count": text.count('\n') + 1
            }

            return text, metadata

        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    text = file.read()
                metadata = {
                    "format": file_ext.lstrip('.') if file_ext else "txt",
                    "size_bytes": os.path.getsize(file_path),
                    "line_count": text.count('\n') + 1,
                    "encoding": "latin-1"
                }
                return text, metadata
            except Exception as e:
                logger.error(f"Error processing file: {str(e)}")
                return "", {"error": f"Processing error: {str(e)}"}
        except Exception as e:
            logger.error(f"Error processing file: {str(e)}")
            return "", {"error": f"Processing error: {str(e)}"}

class SimpleRAGManager:
    """RAG implementation using TF-IDF for first stage and Sentence-BERT for second stage"""
    
    def __init__(self, db_path: str = "rag_documents.db", 
                 use_sbert: bool = True,
                 sbert_model: str = "all-MiniLM-L6-v2",
                 knowledge_base: str = "default"):
        self.db_path = db_path
        self.chunker = DocumentChunker()
        self.knowledge_base = knowledge_base
        
        # Second-stage retrieval with Sentence-BERT (optional)
        self.use_sbert = use_sbert
        self.sbert_model = None
        
        # Knowledge base state cache
        self.kb_states = {}
        
        # Initialize the current state
        self.current_state = self._create_empty_state()
        
        # Only initialize SBERT if requested and available
        if self.use_sbert:
            try:
                from sentence_transformers import SentenceTransformer
                self.sbert_model = SentenceTransformer(sbert_model)
                logger.info(f"Successfully initialized Sentence-BERT model: {sbert_model}")
            except Exception as e:
                logger.error(f"Error initializing Sentence-BERT: {str(e)}")
                logger.warning("Falling back to TF-IDF only mode")
                self.use_sbert = False
        
        # Set up SQLite database for document metadata and content
        self._init_db()
        
        # Load existing documents for the current knowledge base
        # self._load_documents()
        
        class LoadDocumentsWorker(QObject):
            finished = pyqtSignal()
            def __init__(self, manager):
                super().__init__()
                self.manager = manager
            def run(self):
                self.manager._load_documents()
                self.finished.emit()

        self.load_thread = QThread()
        self.load_worker = LoadDocumentsWorker(self)
        self.load_worker.moveToThread(self.load_thread)
        self.load_thread.started.connect(self.load_worker.run)
        self.load_worker.finished.connect(self.load_thread.quit)
        self.load_worker.finished.connect(self.load_worker.deleteLater)
        self.load_thread.finished.connect(self.load_thread.deleteLater)
        self.load_thread.start()
    
    def _create_empty_state(self):
        """Create an empty state for a knowledge base"""
        return {
            "chunks": [],
            "chunk_metadata": [],
            "vectorizer": TfidfVectorizer(),
            "chunk_vectors": None,
            "sbert_vectors": None
        }
    
    def _save_current_state(self):
        """Save the current state to the cache"""
        if self.chunks:
            self.kb_states[self.knowledge_base] = {
                "chunks": self.chunks.copy(),
                "chunk_metadata": self.chunk_metadata.copy(),
                "vectorizer": self.vectorizer,
                "chunk_vectors": self.chunk_vectors,
                "sbert_vectors": self.sbert_vectors
            }
            logger.info(f"Saved state for knowledge base '{self.knowledge_base}'")
    
    def _init_db(self):
        """Initialize SQLite database for document storage"""
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        
        # Create knowledge_bases table if it doesn't exist
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS knowledge_bases (
            id TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            description TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        ''')
        
        # Create documents table if it doesn't exist (updated with knowledge_base_id)
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS documents (
            id TEXT PRIMARY KEY,
            knowledge_base_id TEXT NOT NULL,
            filename TEXT NOT NULL,
            content TEXT NOT NULL,
            metadata TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (knowledge_base_id) REFERENCES knowledge_bases (id)
        )
        ''')
        
        # Create vectors table
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS document_vectors (
            id TEXT PRIMARY KEY,
            document_id TEXT NOT NULL,
            knowledge_base_id TEXT NOT NULL,
            chunk_index INTEGER NOT NULL,
            tfidf_vector BLOB,
            sbert_vector BLOB,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (document_id) REFERENCES documents (id),
            FOREIGN KEY (knowledge_base_id) REFERENCES knowledge_bases (id)
        )
        ''')
        
        # Insert default knowledge base if it doesn't exist
        cursor.execute("SELECT id FROM knowledge_bases WHERE id = 'default'")
        if not cursor.fetchone():
            cursor.execute(
                "INSERT INTO knowledge_bases (id, name, description) VALUES (?, ?, ?)",
                ("default", "None", "Blank knowledge base.")
            )
        
        conn.commit()
        conn.close()
    
    def _compute_sbert_embeddings(self, texts: List[str]) -> np.ndarray:
        """Compute Sentence-BERT embeddings for a list of texts"""
        try:
            # Process in smaller batches if there are many documents
            batch_size = 32
            embeddings = []
            
            for i in range(0, len(texts), batch_size):
                batch = texts[i:i + batch_size]
                batch_embeddings = self.sbert_model.encode(batch, convert_to_numpy=True)
                embeddings.append(batch_embeddings)
            
            if embeddings:
                return np.vstack(embeddings)
            return np.array([])
            
        except Exception as e:
            logger.error(f"Error computing SBERT embeddings: {str(e)}")
            return np.array([])
    
    def _compute_sbert_embeddings_with_progress(self, texts, progress_callback=None):
        """
        Compute Sentence-BERT embeddings with progress reporting
        
        Args:
            texts: List of text chunks
            progress_callback: Optional callback function(current, total)
            
        Returns:
            numpy.ndarray: Embeddings matrix
        """
        try:
            # Process in smaller batches if there are many documents
            batch_size = 32
            embeddings = []
            total = len(texts)
            
            for i in range(0, total, batch_size):
                batch = texts[i:i + batch_size]
                batch_embeddings = self.sbert_model.encode(batch, convert_to_numpy=True)
                embeddings.append(batch_embeddings)
                
                # Report progress if callback provided
                if progress_callback:
                    current = min(i + batch_size, total)
                    progress_callback(current, total)
            
            if embeddings:
                return np.vstack(embeddings)
            return np.array([])
            
        except Exception as e:
            logger.error(f"Error computing SBERT embeddings: {str(e)}")
            return np.array([])
    
    def add_document(self, file_path: str) -> Optional[str]:
        """
        Process and add a document to the RAG system
        
        Args:
            file_path: Path to the document file
            
        Returns:
            doc_id: ID of the added document or None if processing failed
        """
        try:
            # Generate a unique ID for the document
            filename = os.path.basename(file_path)
            doc_id = str(uuid.uuid4())
            
            # Process the document
            content, metadata = DocumentProcessor.process_file(file_path)
            
            if not content:
                logger.error(f"Failed to extract content from {file_path}")
                return None
            
            # Create and chunk the document
            document = Document(
                doc_id=doc_id,
                filename=filename,
                content=content,
                metadata=metadata
            )
            
            chunked_doc = self.chunker.chunk_document(document)
            
            # Store document in SQLite
            self._store_document(chunked_doc)
            
            # Process chunks for vectorization
            new_chunks = []
            new_metadata = []
            
            for i, chunk in enumerate(chunked_doc.chunks):
                new_chunks.append(chunk)
                chunk_meta = {
                    "document_id": document.doc_id,
                    "filename": document.filename,
                    "chunk_index": i,
                    "chunk_count": len(chunked_doc.chunks),
                    "knowledge_base_id": self.knowledge_base
                }
                chunk_meta.update(document.metadata)
                new_metadata.append(chunk_meta)
            
            # Add to existing chunks
            old_chunk_count = len(self.chunks)
            self.chunks.extend(new_chunks)
            self.chunk_metadata.extend(new_metadata)
            
            # Re-vectorize all chunks for TF-IDF
            self.chunk_vectors = self.vectorizer.fit_transform(self.chunks)
            
            # Update SBERT embeddings if enabled
            if self.use_sbert and self.sbert_model is not None:
                try:
                    if self.sbert_vectors is not None and len(self.sbert_vectors) > 0:
                        new_sbert_vectors = self._compute_sbert_embeddings(new_chunks)
                        self.sbert_vectors = np.vstack([self.sbert_vectors, new_sbert_vectors])
                    else:
                        self.sbert_vectors = self._compute_sbert_embeddings(self.chunks)
                except Exception as e:
                    logger.error(f"Error updating SBERT embeddings: {str(e)}")
                    # Continue without SBERT embeddings
                    self.use_sbert = False
            
            # Clear cache for this knowledge base since it's changed
            if self.knowledge_base in self.kb_states:
                del self.kb_states[self.knowledge_base]
                logger.info(f"Cleared cache for knowledge base '{self.knowledge_base}' after document addition")
            
            # After vectorizing the chunks
            if self.chunk_vectors is not None:
                # Get the starting index of the newly added chunks
                start_idx = old_chunk_count
                for i, chunk_idx in enumerate(range(start_idx, len(self.chunks))):
                    doc_id = self.chunk_metadata[chunk_idx]["document_id"]
                    chunk_i = self.chunk_metadata[chunk_idx]["chunk_index"]
                    
                    # Extract TF-IDF vector for this chunk
                    tfidf_vector = self.chunk_vectors[chunk_idx]
                    
                    # Extract SBERT vector if available
                    sbert_vector = None
                    if self.use_sbert and self.sbert_vectors is not None and len(self.sbert_vectors) > chunk_idx:
                        sbert_vector = self.sbert_vectors[chunk_idx]
                    
                    # Store vectors in database
                    self._store_vectors(
                        document_id=doc_id,
                        knowledge_base_id=self.knowledge_base,
                        chunk_index=chunk_i, 
                        tfidf_vector=tfidf_vector,
                        sbert_vector=sbert_vector
                    )
            
            # Return document ID
            return doc_id
            
        except Exception as e:
            logger.error(f"Error adding document: {str(e)}")
            return None
    
    def add_document_async(self, file_path, progress_callback=None, completion_callback=None):
        """
        Process and add a document to the RAG system asynchronously
        
        Args:
            file_path: Path to the document file
            progress_callback: Function to call with progress updates (current, total)
            completion_callback: Function to call when operation completes (success, message)
            
        Returns:
            Worker thread object
        """
        from knowledge_base_worker import KnowledgeBaseWorker
        worker = KnowledgeBaseWorker(self, 'add_document', file_path=file_path)
        
        if progress_callback:
            worker.progress_updated.connect(progress_callback)
        
        if completion_callback:
            worker.operation_completed.connect(completion_callback)
        
        worker.start()
        return worker
        
    def _load_documents(self):
        """Load document metadata and vectors from database for the current knowledge base"""
        try:
            conn = sqlite3.connect(self.db_path)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()
            
            # Clear the current state
            self.current_state["chunks"] = []
            self.current_state["chunk_metadata"] = []
            self.current_state["chunk_vectors"] = None
            self.current_state["sbert_vectors"] = None
            
            # Get metadata for all documents in this knowledge base
            cursor.execute("SELECT id, filename, metadata FROM documents WHERE knowledge_base_id = ?", (self.knowledge_base,))
            doc_meta_rows = cursor.fetchall()
            
            # If we don't have documents, return early
            if not doc_meta_rows:
                logger.info(f"No documents found in knowledge base '{self.knowledge_base}'")
                conn.close()
                return
                
            # Get doc IDs for logging
            doc_ids = [row['id'] for row in doc_meta_rows]
            logger.info(f"Found {len(doc_ids)} documents in knowledge base '{self.knowledge_base}'")
            
            # Now get all document chunks from the vectors table
            # This gives us the chunk structure without having to rechunk the documents
            cursor.execute("""
                SELECT document_id, chunk_index, id 
                FROM document_vectors 
                WHERE knowledge_base_id = ? 
                ORDER BY document_id, chunk_index
            """, (self.knowledge_base,))
            chunk_rows = cursor.fetchall()
            
            # Get the actual document content only if we have no vectors or need to recompute
            if not chunk_rows:
                logger.info(f"No vector chunks found, will need to process documents from scratch")
                
                # Get full document content for processing
                cursor.execute("SELECT id, filename, content, metadata FROM documents WHERE knowledge_base_id = ?", 
                              (self.knowledge_base,))
                doc_rows = cursor.fetchall()
                conn.close()
                
                # Process documents to create chunks
                for doc_data in doc_rows:
                    try:
                        doc_dict = dict(doc_data)
                        metadata = json.loads(doc_data['metadata'] or '{}')
                        
                        document = Document(
                            doc_id=doc_dict['id'],
                            filename=doc_dict['filename'],
                            content=doc_dict['content'],
                            metadata=metadata
                        )
                        
                        # Chunk the document
                        chunked_doc = self.chunker.chunk_document(document)
                        
                        # Add chunks to our collection
                        for i, chunk in enumerate(chunked_doc.chunks):
                            self.current_state["chunks"].append(chunk)
                            chunk_meta = {
                                "document_id": document.doc_id,
                                "filename": document.filename,
                                "chunk_index": i,
                                "chunk_count": len(chunked_doc.chunks),
                                "knowledge_base_id": self.knowledge_base
                            }
                            chunk_meta.update(document.metadata)
                            self.current_state["chunk_metadata"].append(chunk_meta)
                            
                    except Exception as e:
                        logger.error(f"Error processing document {doc_data['id']}: {str(e)}")
                
                # Vectorize all chunks if we have any
                if self.current_state["chunks"]:
                    # First-stage: TF-IDF vectorization
                    self.current_state["chunk_vectors"] = self.current_state["vectorizer"].fit_transform(
                        self.current_state["chunks"]
                    )
                    
                    # Second-stage: Sentence-BERT vectorization (if enabled)
                    if self.use_sbert and self.sbert_model is not None:
                        try:
                            self.current_state["sbert_vectors"] = self._compute_sbert_embeddings(
                                self.current_state["chunks"]
                            )
                            logger.info(f"Generated SBERT embeddings for {len(self.current_state['chunks'])} chunks")
                        except Exception as e:
                            logger.error(f"Error computing SBERT embeddings: {str(e)}")
                            self.use_sbert = False
                    
                    # Save the vectors to the database for future use
                    self._store_all_vectors()
                    
                    logger.info(f"Processed and vectorized {len(self.current_state['chunks'])} chunks from {len(doc_rows)} documents")
                
            else:
                # We have vector records in the database
                # Create a document lookup dictionary for faster access
                doc_lookup = {row['id']: dict(row) for row in doc_meta_rows}
                
                # Build metadata for each chunk based on the vector records
                chunk_ids = []
                for chunk_row in chunk_rows:
                    doc_id = chunk_row['document_id']
                    chunk_index = chunk_row['chunk_index']
                    chunk_id = chunk_row['id']
                    chunk_ids.append(chunk_id)
                    
                    if doc_id in doc_lookup:
                        doc_meta = doc_lookup[doc_id]
                        try:
                            metadata = json.loads(doc_meta['metadata'] or '{}')
                        except (json.JSONDecodeError, TypeError):
                            metadata = {}
                        
                        # Add to chunk metadata
                        chunk_meta = {
                            "document_id": doc_id,
                            "filename": doc_meta['filename'],
                            "chunk_index": chunk_index,
                            "chunk_id": chunk_id,
                            "knowledge_base_id": self.knowledge_base
                        }
                        chunk_meta.update(metadata)
                        self.current_state["chunk_metadata"].append(chunk_meta)
                
                # Close connection after getting metadata
                conn.close()
                
                logger.info(f"Found {len(chunk_ids)} chunk records in the database")
                
                # Now load the actual vectors
                try:
                    # Load vectors from database
                    tfidf_vectors, sbert_vectors, vector_mapping = self._load_vectors(self.knowledge_base)
                    
                    # If we have vectors, use them
                    if tfidf_vectors is not None:
                        self.current_state["chunk_vectors"] = tfidf_vectors
                        logger.info(f"Loaded TF-IDF vectors - shape: {tfidf_vectors.shape}")
                        
                        if sbert_vectors is not None:
                            self.current_state["sbert_vectors"] = sbert_vectors
                            logger.info(f"Loaded SBERT vectors - shape: {sbert_vectors.shape}")
                    
                    # For retrieval, we need the actual chunk content too
                    # But we can load just a placeholder if the vectors are loaded successfully
                    # This will be replaced on demand in retrieve_relevant
                    # This saves memory and processing time
                    placeholder_chunks = [""] * len(self.current_state["chunk_metadata"])
                    self.current_state["chunks"] = placeholder_chunks
                    
                except Exception as e:
                    logger.error(f"Error loading vectors from database: {str(e)}")
                    # Don't try to recover - just return with no vectors
                    return
                    
        except Exception as e:
            logger.error(f"Error in _load_documents: {str(e)}")

    def get_or_compute_vectors(self, kb_id, callback=None):
        """
        Get vectors from cache or database, computing only if necessary
        
        Args:
            kb_id: Knowledge base ID
            callback: Function to call when operation completes (success, message)
            
        Returns:
            (tfidf_vectors, sbert_vectors) tuple or None if async
        """
        # First check the cache
        if kb_id in self.kb_states:
            state = self.kb_states[kb_id]
            if state["chunk_vectors"] is not None:
                logger.info(f"Using cached vectors for knowledge base '{kb_id}'")
                if callback:
                    callback(True, "Using cached vectors")
                return state["chunk_vectors"], state["sbert_vectors"]
        
        # Then try to load from database
        from knowledge_base_worker import KnowledgeBaseWorker
        worker = KnowledgeBaseWorker(self, 'load_vectors', kb_id=kb_id)
        
        def on_vectors_loaded(tfidf_vectors, sbert_vectors, vector_mapping):
            # If we have vectors in the database, use them
            if tfidf_vectors is not None and len(tfidf_vectors) > 0:
                self.current_state["chunk_vectors"] = tfidf_vectors
                logger.info(f"Loaded TF-IDF vectors from database for knowledge base '{kb_id}'")
                
                if sbert_vectors is not None and len(sbert_vectors) > 0:
                    self.current_state["sbert_vectors"] = sbert_vectors
                    logger.info(f"Loaded SBERT vectors from database for knowledge base '{kb_id}'")
                
                # Update the cache
                if kb_id in self.kb_states:
                    self.kb_states[kb_id]["chunk_vectors"] = tfidf_vectors
                    self.kb_states[kb_id]["sbert_vectors"] = sbert_vectors
                
                if callback:
                    callback(True, "Loaded vectors from database")
            else:
                # Need to compute vectors
                self._compute_vectors_async(callback)
        
        worker.vectors_loaded.connect(on_vectors_loaded)
        worker.start()
        return None, None

    def _compute_vectors_async(self, callback=None):
        """Compute vectors asynchronously"""
        # Vectorize all chunks if we have any
        if not self.current_state["chunks"]:
            if callback:
                callback(False, "No chunks to vectorize")
            return
        
        # First-stage: TF-IDF vectorization (this is fast, do it synchronously)
        try:
            self.current_state["chunk_vectors"] = self.current_state["vectorizer"].fit_transform(
                self.current_state["chunks"]
            )
            
            # Ensure the chunk_vectors is properly initialized
            if self.current_state["chunk_vectors"] is None or self.current_state["chunk_vectors"].shape[0] != len(self.current_state["chunks"]):
                logger.error(f"Error in TF-IDF vectorization: shape mismatch - got {self.current_state['chunk_vectors'].shape if self.current_state['chunk_vectors'] is not None else None}, expected ({len(self.current_state['chunks'])}, N)")
                if callback:
                    callback(False, "Error in TF-IDF vectorization: shape mismatch")
                return
        except Exception as e:
            logger.error(f"Error computing TF-IDF vectors: {str(e)}")
            if callback:
                callback(False, f"Error computing TF-IDF vectors: {str(e)}")
            return
        
        # Second-stage: Sentence-BERT vectorization (this is slow, do it asynchronously)
        if self.use_sbert and self.sbert_model is not None:
            try:
                from knowledge_base_worker import KnowledgeBaseWorker
                worker = KnowledgeBaseWorker(
                    self, 
                    'compute_embeddings', 
                    chunks=self.current_state["chunks"]
                )
                
                def on_embeddings_completed(success, message):
                    if success:
                        # Store the computed embeddings
                        if hasattr(worker, 'result') and worker.result is not None:
                            self.current_state["sbert_vectors"] = worker.result
                            
                            # Save to cache
                            if self.knowledge_base in self.kb_states:
                                self.kb_states[self.knowledge_base]["sbert_vectors"] = worker.result
                        
                        logger.info(f"Generated SBERT embeddings for {len(self.current_state['chunks'])} chunks")
                        
                        # Store vectors in database for future use
                        self._store_all_vectors()
                    
                    if callback:
                        callback(success, message)
                
                worker.operation_completed.connect(on_embeddings_completed)
                worker.start()
                
            except Exception as e:
                logger.error(f"Error computing SBERT embeddings asynchronously: {str(e)}")
                self.use_sbert = False
                if callback:
                    callback(False, f"Error computing embeddings: {str(e)}")
        else:
            # No SBERT, just use TF-IDF
            # Store the TF-IDF vectors in the database
            self._store_all_vectors()
            
            if callback:
                callback(True, "TF-IDF vectors computed (SBERT not enabled)")

    def _store_all_vectors(self):
        """Store all vectors for the current knowledge base in the database"""
        if not self.current_state["chunk_vectors"] or not self.current_state["chunk_metadata"]:
            logger.warning("No vectors or metadata to store")
            return
        
        try:    
            # Check if the number of vectors matches the number of chunks
            vector_count = self.current_state["chunk_vectors"].shape[0]
            metadata_count = len(self.current_state["chunk_metadata"])
            
            if vector_count != metadata_count:
                logger.error(f"Vector count ({vector_count}) does not match metadata count ({metadata_count})")
                return
                
            for i, meta in enumerate(self.current_state["chunk_metadata"]):
                try:
                    # Extract TF-IDF vector for this chunk
                    tfidf_vector = self.current_state["chunk_vectors"][i]
                    
                    # Extract SBERT vector if available
                    sbert_vector = None
                    if (self.use_sbert and 
                        self.current_state["sbert_vectors"] is not None and 
                        i < len(self.current_state["sbert_vectors"])):
                        sbert_vector = self.current_state["sbert_vectors"][i]
                    
                    # Store vectors in database
                    self._store_vectors(
                        document_id=meta["document_id"],
                        knowledge_base_id=self.knowledge_base,
                        chunk_index=meta["chunk_index"], 
                        tfidf_vector=tfidf_vector,
                        sbert_vector=sbert_vector
                    )
                except Exception as e:
                    logger.error(f"Error storing vector for chunk {i}: {str(e)}")
            
            logger.info(f"Stored vectors for {len(self.current_state['chunk_metadata'])} chunks in knowledge base '{self.knowledge_base}'")
        except Exception as e:
            logger.error(f"Error in _store_all_vectors: {str(e)}")
    
    def retrieve_relevant(self, query: str, top_k_tfidf: int = 10, top_k_final: int = 3) -> List[Dict[str, Any]]:
        """
        Two-stage retrieval: first TF-IDF, then re-rank with Sentence-BERT if available
        
        Args:
            query: Query text
            top_k_tfidf: Number of documents to retrieve with TF-IDF (default 10)
            top_k_final: Number of final results after SBERT re-ranking (default 3)
            
        Returns:
            List of relevant chunks with metadata
        """
        try:
            # Check if we have vectors
            if self.chunk_vectors is None or self.chunk_vectors.shape[0] == 0 or len(self.chunk_metadata) == 0:
                logger.warning("No vectors available for retrieval")
                return []

            
            # Check if chunks are placeholders (empty strings) and load content if needed
            # Check if chunks are placeholders (empty strings) and load content for retrieval
            if self.chunks and all(not chunk for chunk in self.chunks):
                logger.info("Chunks are placeholders, loading content for retrieval...")
                loaded_chunks = self._load_chunk_content_for_retrieval()
                if loaded_chunks:
                    self.chunks = loaded_chunks
                    logger.info(f"Loaded {len(loaded_chunks)} chunks for retrieval")
                    # Fit the TF-IDF vectorizer on the loaded chunks.
                    logger.info("Fitting TF-IDF vectorizer on loaded chunks")
                    self.vectorizer.fit(self.chunks)
                    # Update the TF-IDF vectors so that they match the new vocabulary.
                    self.chunk_vectors = self.vectorizer.transform(self.chunks)


            
            # Make sure we have valid chunks
            if not self.chunks or len(self.chunks) == 0 or len(self.chunks) != len(self.chunk_metadata):
                logger.warning(f"Chunk inconsistency: {len(self.chunks)} chunks vs {len(self.chunk_metadata)} metadata")
                return []
            
            # Stage 1: TF-IDF retrieval
            logger.info(f"Performing TF-IDF retrieval with {top_k_tfidf} candidates...")
            query_vector = self.vectorizer.transform([query])
            similarities = cosine_similarity(query_vector, self.chunk_vectors).flatten()
            
            # Get top-k indices from TF-IDF
            candidate_count = min(top_k_tfidf, len(self.chunks))
            candidate_indices = similarities.argsort()[-candidate_count:][::-1]
            
            # If SBERT is not available or disabled, return TF-IDF results
            if not self.use_sbert or self.sbert_model is None or self.sbert_vectors is None:
                # Format TF-IDF results
                results = []
                for i, idx in enumerate(candidate_indices[:top_k_final]):
                    if idx >= len(self.chunks) or idx >= len(self.chunk_metadata):
                        logger.warning(f"Index out of bounds: {idx} for chunks length {len(self.chunks)}")
                        continue
                        
                    chunk = self.chunks[idx]
                    metadata = self.chunk_metadata[idx]
                    
                    results.append({
                        "chunk_id": f"{metadata['document_id']}_chunk_{metadata['chunk_index']}",
                        "document_id": metadata['document_id'],
                        "filename": metadata['filename'],
                        "chunk_index": metadata['chunk_index'],
                        "content": chunk,
                        "similarity": float(similarities[idx]),
                        "method": "tfidf",
                        "metadata": metadata
                    })
                
                logger.info(f"Returning {len(results)} results using TF-IDF only")
                return results
            
            # Stage 2: Re-rank with Sentence-BERT
            try:
                logger.info("Performing SBERT re-ranking...")
                candidate_chunks = []
                candidate_meta = []
                valid_indices = []
                
                # Validate candidates
                for i, idx in enumerate(candidate_indices):
                    if idx < len(self.chunks) and idx < len(self.sbert_vectors):
                        candidate_chunks.append(self.chunks[idx])
                        candidate_meta.append(self.chunk_metadata[idx])
                        valid_indices.append(idx)
                    else:
                        logger.warning(f"Invalid candidate index: {idx}")
                
                if not candidate_chunks:
                    logger.warning("No valid candidates for SBERT re-ranking")
                    return []
                    
                # Get embeddings for candidates
                candidate_embeddings = np.array([self.sbert_vectors[idx] for idx in valid_indices])
                
                # Get query embedding with SBERT
                query_embedding = self.sbert_model.encode([query], convert_to_numpy=True)
                
                # Calculate cosine similarities with SBERT
                sbert_similarities = cosine_similarity(query_embedding, candidate_embeddings).flatten()
                
                # Get top-k indices after re-ranking
                top_k = min(top_k_final, len(valid_indices))
                final_indices = sbert_similarities.argsort()[-top_k:][::-1]
                
                # Format results
                results = []
                for i, rel_idx in enumerate(final_indices):
                    chunk = candidate_chunks[rel_idx]
                    metadata = candidate_meta[rel_idx]
                    abs_idx = valid_indices[rel_idx]
                    
                    results.append({
                        "chunk_id": f"{metadata['document_id']}_chunk_{metadata['chunk_index']}",
                        "document_id": metadata['document_id'],
                        "filename": metadata['filename'],
                        "chunk_index": metadata['chunk_index'],
                        "content": chunk,
                        "tfidf_similarity": float(similarities[abs_idx]),
                        "sbert_similarity": float(sbert_similarities[rel_idx]),
                        "similarity": float(sbert_similarities[rel_idx]),  # Use SBERT score as primary
                        "method": "sbert",
                        "metadata": metadata
                    })
                
                logger.info(f"Returning {len(results)} results using SBERT re-ranking")
                return results
                
            except Exception as e:
                logger.error(f"Error during SBERT re-ranking: {str(e)}")
                # Fallback to TF-IDF results if SBERT fails
                results = []
                for i, idx in enumerate(candidate_indices[:top_k_final]):
                    if idx >= len(self.chunks) or idx >= len(self.chunk_metadata):
                        continue
                        
                    chunk = self.chunks[idx]
                    metadata = self.chunk_metadata[idx]
                    
                    results.append({
                        "chunk_id": f"{metadata['document_id']}_chunk_{metadata['chunk_index']}",
                        "document_id": metadata['document_id'],
                        "filename": metadata['filename'],
                        "chunk_index": metadata['chunk_index'],
                        "content": chunk,
                        "similarity": float(similarities[idx]),
                        "method": "tfidf_fallback",
                        "metadata": metadata
                    })
                
                logger.info(f"Returning {len(results)} results using TF-IDF fallback")
                return results
                
        except Exception as e:
            logger.error(f"Error retrieving documents: {str(e)}")
            return []

    def _load_chunk_content_for_retrieval(self):
        """Load actual chunk content for placeholder chunks"""
        try:
            # Get a list of document IDs needed
            doc_ids = set(meta['document_id'] for meta in self.chunk_metadata)
            
            # Load documents
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            placeholders = ', '.join(['?'] * len(doc_ids))
            query = f"SELECT id, content FROM documents WHERE id IN ({placeholders})"
            cursor.execute(query, list(doc_ids))
            
            doc_contents = {row[0]: row[1] for row in cursor.fetchall()}
            conn.close()
            
            # If we couldn't load any documents, return
            if not doc_contents:
                logger.warning("Could not load document content for chunks")
                return None
            
            # Rechunk documents to get actual chunk content
            chunks = []
            for meta in self.chunk_metadata:
                doc_id = meta['document_id']
                chunk_index = meta['chunk_index']
                
                if doc_id in doc_contents:
                    # Create document object
                    doc = Document(
                        doc_id=doc_id,
                        filename=meta['filename'],
                        content=doc_contents[doc_id]
                    )
                    
                    # Chunk the document
                    chunked_doc = self.chunker.chunk_document(doc)
                    
                    # Get the specific chunk we need
                    if chunk_index < len(chunked_doc.chunks):
                        chunks.append(chunked_doc.chunks[chunk_index])
                    else:
                        logger.warning(f"Chunk index {chunk_index} out of bounds for document {doc_id}")
                        chunks.append("")  # Add empty chunk as placeholder
                else:
                    logger.warning(f"Document content not found for doc_id {doc_id}")
                    chunks.append("")  # Add empty chunk as placeholder
            
            return chunks
            
        except Exception as e:
            logger.error(f"Error loading chunk content: {str(e)}")
            return None
            
    def _store_document(self, document: Document):
        """Store document in SQLite database"""
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        
        cursor.execute(
            "INSERT INTO documents (id, knowledge_base_id, filename, content, metadata) VALUES (?, ?, ?, ?, ?)",
            (
                document.doc_id,
                self.knowledge_base,
                document.filename,
                document.content,
                json.dumps(document.metadata)
            )
        )
        
        conn.commit()
        conn.close()
    
    def get_document(self, doc_id: str) -> Optional[Document]:
        """Get a document by ID"""
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute("SELECT * FROM documents WHERE id = ?", (doc_id,))
        row = cursor.fetchone()
        
        conn.close()
        
        if row:
            # Convert row to dict
            doc_dict = dict(row)
            
            # Parse metadata from JSON
            try:
                metadata = json.loads(doc_dict["metadata"] or "{}")
            except (json.JSONDecodeError, TypeError):
                metadata = {}
            
            return Document(
                doc_id=doc_dict["id"],
                filename=doc_dict["filename"],
                content=doc_dict["content"],
                metadata=metadata
            )
        
        return None

    def get_all_documents(self) -> List[Dict[str, Any]]:
        """Get all documents' metadata (without content) for the current knowledge base"""
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute(
            "SELECT id, filename, metadata, created_at FROM documents WHERE knowledge_base_id = ?", 
            (self.knowledge_base,)
        )
        rows = cursor.fetchall()
        
        conn.close()
        
        documents = []
        for row in rows:
            row_dict = dict(row)
            try:
                metadata = json.loads(row_dict["metadata"] or "{}")
            except (json.JSONDecodeError, TypeError):
                metadata = {}
            
            documents.append({
                "id": row_dict["id"],
                "filename": row_dict["filename"],
                "metadata": metadata,
                "created_at": row_dict["created_at"]
            })
        
        return documents
    
    def delete_document(self, doc_id: str) -> bool:
        """Delete a document and its vectors"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            # First delete the vectors
            cursor.execute("DELETE FROM document_vectors WHERE document_id = ?", (doc_id,))
            
            # Then delete the document
            cursor.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
            deleted = cursor.rowcount > 0
            
            conn.commit()
            conn.close()
            
            if deleted:
                # Clear cache for this knowledge base
                if self.knowledge_base in self.kb_states:
                    del self.kb_states[self.knowledge_base]
                    logger.info(f"Cleared cache for knowledge base '{self.knowledge_base}' after document deletion")
                
                # Reload documents
                self.current_state = self._create_empty_state()
                self._load_documents()
            
            return deleted
            
        except Exception as e:
            logger.error(f"Error deleting document: {str(e)}")
            return False
    
 

    # Add methods to store and retrieve vectors
    def _store_vectors(self, document_id, knowledge_base_id, chunk_index, tfidf_vector, sbert_vector=None):
        """Store document vectors in the database"""
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        
        vector_id = f"{document_id}_chunk_{chunk_index}"
        
        # Serialize vectors
        tfidf_blob = pickle.dumps(tfidf_vector)
        sbert_blob = pickle.dumps(sbert_vector) if sbert_vector is not None else None
        
        # Use REPLACE to handle updates
        cursor.execute(
            """REPLACE INTO document_vectors 
               (id, document_id, knowledge_base_id, chunk_index, tfidf_vector, sbert_vector)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (vector_id, document_id, knowledge_base_id, chunk_index, tfidf_blob, sbert_blob)
        )
        
        conn.commit()
        conn.close()
        
    def _load_vectors(self, knowledge_base_id):
        """Load vectors for a knowledge base from the database"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            # Get all vectors for this knowledge base
            cursor.execute(
                "SELECT document_id, chunk_index, tfidf_vector, sbert_vector FROM document_vectors WHERE knowledge_base_id = ?",
                (knowledge_base_id,)
            )
            
            vector_rows = cursor.fetchall()
            conn.close()
            
            # Check if we found anything
            if not vector_rows:
                logger.info(f"No vectors found in database for knowledge base '{knowledge_base_id}'")
                return None, None, {}
            
            logger.info(f"Found {len(vector_rows)} vector records in the database")
            
            # Initialize vectors and mapping
            tfidf_vectors = []
            sbert_vectors = []
            vector_mapping = {}
            
            # Count errors for logging
            tfidf_errors = 0
            sbert_errors = 0
            
            # Process each vector row
            for i, (doc_id, chunk_idx, tfidf_blob, sbert_blob) in enumerate(vector_rows):
                # Process TFIDF vector
                try:
                    if tfidf_blob is not None:
                        tfidf_vector = pickle.loads(tfidf_blob)
                        tfidf_vectors.append(tfidf_vector)
                        vector_mapping[(doc_id, chunk_idx)] = len(tfidf_vectors) - 1
                        
                        # Process SBERT vector if available
                        if sbert_blob is not None:
                            try:
                                sbert_vector = pickle.loads(sbert_blob)
                                sbert_vectors.append(sbert_vector)
                            except Exception as e:
                                sbert_errors += 1
                except Exception as e:
                    tfidf_errors += 1
            
            if tfidf_errors > 0:
                logger.warning(f"Encountered {tfidf_errors} errors loading TF-IDF vectors")
            if sbert_errors > 0:
                logger.warning(f"Encountered {sbert_errors} errors loading SBERT vectors")
                
            # Check if we loaded any vectors
            if not tfidf_vectors:
                logger.warning(f"No valid TF-IDF vectors could be loaded for knowledge base '{knowledge_base_id}'")
                return None, None, {}
            
            # Create combined vectors
            try:
                # For TF-IDF vectors (sparse)
                combined_tfidf = scipy.sparse.vstack(tfidf_vectors)
                logger.info(f"Successfully created combined TF-IDF matrix with shape {combined_tfidf.shape}")
                
                # For SBERT vectors (dense)
                combined_sbert = None
                if sbert_vectors:
                    combined_sbert = np.vstack(sbert_vectors)
                    logger.info(f"Successfully created combined SBERT matrix with shape {combined_sbert.shape}")
                    
                return combined_tfidf, combined_sbert, vector_mapping
                
            except Exception as e:
                logger.error(f"Error combining vectors: {str(e)}")
                
                # Last resort fallback - try with error correction
                try:
                    logger.info("Attempting fallback vector stacking...")
                    
                    # Get the first vector to determine expected shape
                    ref_vector = tfidf_vectors[0]
                    if scipy.sparse.issparse(ref_vector):
                        # For sparse vectors
                        n_features = ref_vector.shape[1]
                        
                        # Create an empty matrix and fill it
                        combined_matrix = scipy.sparse.csr_matrix((len(tfidf_vectors), n_features))
                        for i, vec in enumerate(tfidf_vectors):
                            if vec.shape[1] == n_features:
                                combined_matrix[i] = vec
                        
                        logger.info(f"Fallback method created TF-IDF matrix with shape {combined_matrix.shape}")
                        return combined_matrix, None, vector_mapping
                    else:
                        # For dense vectors (unlikely for TF-IDF)
                        return None, None, {}
                except Exception as e2:
                    logger.error(f"Fallback vector stacking also failed: {str(e2)}")
                    return None, None, {}
                    
        except Exception as e:
            logger.error(f"Error in _load_vectors: {str(e)}")
            return None, None, {}
    
    def set_knowledge_base(self, knowledge_base_id: str):
        """Switch to a different knowledge base"""
        # If it's the same knowledge base, do nothing
        if self.knowledge_base == knowledge_base_id:
            return
        
        # Save current state to cache
        self._save_current_state()
        
        # Set new knowledge base
        self.knowledge_base = knowledge_base_id
        
        # Check if we have cached state for this knowledge base
        if knowledge_base_id in self.kb_states:
            logger.info(f"Loading cached state for knowledge base '{knowledge_base_id}'")
            self.current_state = self.kb_states[knowledge_base_id]
        else:
            # Create a new state for this knowledge base
            logger.info(f"No cache found for knowledge base '{knowledge_base_id}', loading from database")
            self.current_state = self._create_empty_state()
            self._load_documents()
    
    def set_knowledge_base_async(self, knowledge_base_id, callback=None):
        """
        Asynchronously switch to a different knowledge base
        
        Args:
            knowledge_base_id: ID of the knowledge base to switch to
            callback: Function to call when operation completes (success, message)
        """
        # If it's the same knowledge base, do nothing
        if self.knowledge_base == knowledge_base_id:
            if callback:
                callback(True, "Already using this knowledge base")
            return
        
        # Save current state to cache
        self._save_current_state()
        
        # Set new knowledge base
        self.knowledge_base = knowledge_base_id
        
        # Check if we have cached state for this knowledge base
        if knowledge_base_id in self.kb_states:
            logger.info(f"Loading cached state for knowledge base '{knowledge_base_id}'")
            self.current_state = self.kb_states[knowledge_base_id]
            if callback:
                callback(True, "Loaded knowledge base from cache")
        else:
            # Create a new state for this knowledge base
            logger.info(f"No cache found for knowledge base '{knowledge_base_id}', loading from database")
            self.current_state = self._create_empty_state()
            
            # Start asynchronous loading
            from knowledge_base_worker import KnowledgeBaseWorker
            self.kb_worker = KnowledgeBaseWorker(self, 'load_knowledge_base', kb_id=knowledge_base_id)
            
            if callback:
                self.kb_worker.operation_completed.connect(callback)
            
            self.kb_worker.start()
    
    def get_knowledge_bases(self) -> List[Dict[str, Any]]:
        """Get all knowledge bases"""
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute("SELECT * FROM knowledge_bases ORDER BY name")
        rows = cursor.fetchall()
        
        conn.close()
        
        knowledge_bases = []
        for row in rows:
            knowledge_bases.append(dict(row))
        
        return knowledge_bases

    def create_knowledge_base(self, name: str, description: str = "") -> str:
        """Create a new knowledge base"""
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        
        kb_id = str(uuid.uuid4())
        
        cursor.execute(
            "INSERT INTO knowledge_bases (id, name, description) VALUES (?, ?, ?)",
            (kb_id, name, description)
        )
        
        conn.commit()
        conn.close()
        
        return kb_id
        
    def delete_knowledge_base(self, kb_id: str) -> bool:
        """Delete a knowledge base and all its documents"""
        try:
            # Don't allow deleting the default knowledge base
            if kb_id == "default":
                logger.warning("Cannot delete the default knowledge base")
                return False
                
            # Delete from SQLite - first documents, then knowledge base
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            # Delete all documents in this knowledge base
            cursor.execute("DELETE FROM documents WHERE knowledge_base_id = ?", (kb_id,))
            
            # Delete the knowledge base itself
            cursor.execute("DELETE FROM knowledge_bases WHERE id = ?", (kb_id,))
            deleted = cursor.rowcount > 0
            
            conn.commit()
            conn.close()
            
            # Clear cache for this knowledge base
            if kb_id in self.kb_states:
                del self.kb_states[kb_id]
                logger.info(f"Cleared cache for knowledge base '{kb_id}' after deletion")
            
            # If we deleted the current knowledge base, switch to default
            if self.knowledge_base == kb_id:
                self.set_knowledge_base("default")
            
            return deleted
            
        except Exception as e:
            logger.error(f"Error deleting knowledge base: {str(e)}")
            return False
            
    def save_embeddings(self, path: str = "embeddings.npz"):
        """Save vector embeddings to disk to avoid recomputation"""
        try:
            if self.sbert_vectors is not None and len(self.sbert_vectors) > 0:
                np.savez_compressed(
                    path, 
                    sbert_vectors=self.sbert_vectors,
                    chunk_ids=[f"{m['document_id']}_{m['chunk_index']}" for m in self.chunk_metadata]
                )
                logger.info(f"Saved embeddings to {path}")
                return True
            return False
        except Exception as e:
            logger.error(f"Error saving embeddings: {str(e)}")
            return False

    def load_embeddings(self, path: str = "embeddings.npz"):
        """Load vector embeddings from disk"""
        try:
            if not os.path.exists(path):
                logger.warning(f"Embeddings file {path} not found")
                return False
                
            data = np.load(path)
            saved_chunk_ids = data['chunk_ids']
            current_chunk_ids = [f"{m['document_id']}_{m['chunk_index']}" for m in self.chunk_metadata]
            
            # Only load if the chunks match exactly
            if len(saved_chunk_ids) == len(current_chunk_ids) and set(saved_chunk_ids) == set(current_chunk_ids):
                self.sbert_vectors = data['sbert_vectors']
                logger.info(f"Loaded embeddings from {path}")
                return True
            else:
                logger.warning("Saved embeddings don't match current chunks. Recomputing...")
                self.sbert_vectors = self._compute_sbert_embeddings(self.chunks)
                return False
        except Exception as e:
            logger.error(f"Error loading embeddings: {str(e)}")
            return False
    
    def terminate_workers(self):
        """Terminate all running worker threads"""
        if hasattr(self, 'kb_worker') and self.kb_worker.isRunning():
            self.kb_worker.terminate()
            self.kb_worker.wait()
            
    def __del__(self):
        """Cleanup when object is destroyed"""
        # Ensure all worker threads are stopped
        for attr_name in dir(self):
            attr = getattr(self, attr_name)
            if isinstance(attr, QThread) and attr.isRunning():
                try:
                    attr.terminate()
                    attr.wait()
                except (RuntimeError, Exception) as e:
                    logger.warning(f"Error terminating worker: {e}")

    # Property getters and setters for the current state
    @property
    def chunks(self):
        return self.current_state["chunks"]
    
    @chunks.setter
    def chunks(self, value):
        self.current_state["chunks"] = value
    
    @property
    def chunk_metadata(self):
        return self.current_state["chunk_metadata"]
    
    @chunk_metadata.setter
    def chunk_metadata(self, value):
        self.current_state["chunk_metadata"] = value
    
    @property
    def vectorizer(self):
        return self.current_state["vectorizer"]
    
    @vectorizer.setter
    def vectorizer(self, value):
        self.current_state["vectorizer"] = value
    
    @property
    def chunk_vectors(self):
        return self.current_state["chunk_vectors"]
    
    @chunk_vectors.setter
    def chunk_vectors(self, value):
        self.current_state["chunk_vectors"] = value
    
    @property
    def sbert_vectors(self):
        return self.current_state["sbert_vectors"]
    
    @sbert_vectors.setter
    def sbert_vectors(self, value):
        self.current_state["sbert_vectors"] = value